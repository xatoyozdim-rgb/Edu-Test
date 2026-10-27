import io
from docx import Document
from django.conf import settings
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth import authenticate, login, logout
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.http import JsonResponse, HttpResponseForbidden, HttpResponse
from django.utils import timezone
from django.utils.dateparse import parse_datetime
from django.db.models import Avg, Max, Count, Q
from .models import User, Test, Question, Option, TestAttempt, StudentAnswer, Complaint
from .utils import parse_docx_to_test, recalculate_test_scores
from .google_cloud import (
    get_google_auth_url, exchange_code_for_tokens, get_google_user_info,
    sync_test_to_teacher_google_cloud, serialize_test_payload
)

def login_view(request):
    if request.user.is_authenticated:
        return redirect('dashboard')
    if request.method == 'POST':
        un = request.POST.get('username', '').strip()
        pw = request.POST.get('password', '')
        # Ustozlar oddiy login/parol bilan kira olmaydi — faqat Google akkaunt orqali!
        existing_user = User.objects.filter(username=un).first()
        if existing_user and existing_user.is_teacher():
            messages.error(request, "Ustozlar tizimga faqat Google akkaunti orqali kirishi shart! Quyidagi 'Google orqali kirish' tugmasini bosing.")
            return render(request, 'accounts/login.html')

        u = authenticate(request, username=un, password=pw)
        if u:
            if u.is_teacher():
                messages.error(request, "Ustozlar faqat Google akkaunti orqali tizimga kirishi mumkin!")
                return render(request, 'accounts/login.html')
            login(request, u)
            return redirect('dashboard')
        messages.error(request, "Login yoki parol xato!")
    return render(request, 'accounts/login.html')

def signup_view(request):
    if request.user.is_authenticated:
        return redirect('dashboard')
    if request.method == 'POST':
        role = request.POST.get('role', 'student')
        fn = request.POST.get('full_name', '').strip()
        subject = request.POST.get('subject', '').strip()
        cloud_bucket = request.POST.get('google_cloud_bucket', '').strip()

        # 3-talab: Ustoz ro'yxatdan o'tayotganda ma'lumotlar to'ldirilgach faqat Google akkaunt orqali kiradi
        if role == 'teacher':
            if not fn:
                messages.error(request, "Ism va familiyangizni kiriting!")
                return render(request, 'accounts/signup.html')
            request.session['pending_teacher_signup'] = {
                'full_name': fn,
                'subject': subject,
                'google_cloud_bucket': cloud_bucket,
            }
            return redirect('google_login')

        # O'quvchilar uchun odatiy ro'yxatdan o'tish
        un = request.POST.get('username', '').strip()
        pw = request.POST.get('password', '').strip()
        if not un or not pw:
            messages.error(request, "O'quvchi uchun login va parol kiritilishi shart!")
            return render(request, 'accounts/signup.html')
        if not User.objects.filter(username=un).exists():
            u = User.objects.create_user(
                username=un, password=pw, full_name=fn,
                role='student', grade=request.POST.get('grade', 'barcha'),
                letter=request.POST.get('letter', 'barchasi')
            )
            login(request, u)
            return redirect('dashboard')
        messages.error(request, "Bu login band!")
    return render(request, 'accounts/signup.html')

def google_login_start(request):
    """Ustozlarni Google OAuth 2.0 sahifasiga yo'naltiradi."""
    if settings.GOOGLE_CLIENT_ID and settings.GOOGLE_CLIENT_SECRET:
        return redirect(get_google_auth_url(state="teacher_auth"))
    # Agar GOOGLE_CLIENT_ID hali sozlanmagan bo'lsa, Google OAuth interfeys oynasini ochadi
    pending = request.session.get('pending_teacher_signup', {})
    return render(request, 'accounts/google_oauth_consent.html', {'pending': pending})

def google_login_callback(request):
    """Google OAuth 2.0 qaytimini (yoki lokal Google tasdiqlashni) qabul qiladi va ustozni tizimga kiritadi."""
    pending = request.session.pop('pending_teacher_signup', {})

    if request.method == 'POST' and not settings.GOOGLE_CLIENT_ID:
        # Lokal/Dev rejimdagi Google Account tasdiqlash
        g_email = request.POST.get('google_email', '').strip().lower()
        if not g_email or '@' not in g_email:
            messages.error(request, "To'g'ri Google (Gmail) manzilini kiriting!")
            return redirect('google_login')
        g_sub = f"google_uid_{g_email}"
        g_name = pending.get('full_name') or request.POST.get('google_name', '').strip() or g_email.split('@')[0].title()
        access_token = f"dev-token-{g_email}"
        refresh_token = f"dev-refresh-{g_email}"
    else:
        code = request.GET.get('code')
        if not code:
            messages.error(request, "Google avtorizatsiya bekor qilindi.")
            return redirect('login')
        try:
            token_data = exchange_code_for_tokens(code)
            access_token = token_data.get('access_token')
            refresh_token = token_data.get('refresh_token', '')
            info = get_google_user_info(access_token)
            g_email = info.get('email', '').lower()
            g_sub = info.get('sub')
            g_name = pending.get('full_name') or info.get('name') or g_email.split('@')[0]
        except Exception as e:
            messages.error(request, f"Google orqali kirishda xatolik: {e}")
            return redirect('login')

    # Ustozni Google ID yoki Google Email orqali topamiz yoki yangi yaratamiz
    user = User.objects.filter(Q(google_id=g_sub) | Q(google_email=g_email) | Q(username=g_email)).first()
    if not user:
        base_username = g_email.split('@')[0]
        uname = base_username
        counter = 1
        while User.objects.filter(username=uname).exists():
            uname = f"{base_username}_{counter}"
            counter += 1
        user = User(
            username=uname,
            email=g_email,
            google_email=g_email,
            google_id=g_sub,
            full_name=g_name,
            role='teacher',
            subject=pending.get('subject', '') or request.POST.get('subject', ''),
            google_cloud_bucket=pending.get('google_cloud_bucket', '') or request.POST.get('google_cloud_bucket', ''),
        )
        # Parol orqali kirishni butunlay yopamiz (faqat Google orqali kiradi)
        user.set_unusable_password()
        user.save()
    else:
        user.role = 'teacher'
        user.google_id = g_sub
        user.google_email = g_email
        user.email = g_email
        if pending.get('full_name'):
            user.full_name = pending['full_name']
        if pending.get('subject'):
            user.subject = pending['subject']
        if pending.get('google_cloud_bucket'):
            user.google_cloud_bucket = pending['google_cloud_bucket']
        user.set_unusable_password()

    user.google_access_token = access_token
    if refresh_token:
        user.google_refresh_token = refresh_token
    user.save()

    login(request, user)
    messages.success(request, f"Xush kelibsiz, {user.get_display_name()}! Shaxsiy Google Cloud ({user.google_email}) faollashtirildi.")
    return redirect('dashboard')

@login_required
def profile_view(request):
    if request.method == 'POST':
        request.user.full_name = request.POST.get('full_name', '').strip()
        request.user.subject = request.POST.get('subject', '').strip()
        request.user.grade = request.POST.get('grade', 'barcha')
        request.user.letter = request.POST.get('letter', 'barchasi')
        if request.user.is_teacher():
            request.user.google_cloud_bucket = request.POST.get('google_cloud_bucket', '').strip()
        else:
            pwd = request.POST.get('password', '').strip()
            if pwd:
                request.user.set_password(pwd)
        request.user.save()
        if not request.user.is_teacher() and request.POST.get('password', '').strip():
            login(request, request.user)
        messages.success(request, "Ma'lumotlaringiz muvaffaqiyatli yangilandi!")
        return redirect('profile')
    return render(request, 'accounts/profile.html')

def logout_view(request):
    logout(request)
    return redirect('login')

@login_required
def dashboard_view(request):
    now = timezone.now()

    if request.user.is_teacher():
        tests = Test.objects.filter(creator=request.user).annotate(attempt_count=Count('attempts'), avg_score=Avg('attempts__score'))
        
        teachers_query = User.objects.filter(role='teacher')
        raw_leaderboard = []
        for th in teachers_query:
            th_tests = Test.objects.filter(creator=th)
            th_test_count = th_tests.count()
            th_solves = TestAttempt.objects.filter(test__creator=th, status='completed').count()
            ratio = (th_solves / th_test_count) if th_test_count > 0 else 0.0
            raw_leaderboard.append({'teacher': th, 'test_count': th_test_count, 'total_solves': th_solves, 'ratio': ratio})
        
        raw_leaderboard = sorted(raw_leaderboard, key=lambda x: x['total_solves'], reverse=True)

        # 1-talab: Ustozlar reytingida effektlar:
        # 1, 2, 3-o'rin uchun bir xil tilla rang effekt;
        # Top 10% yashil effekt;
        # Top 70% ko'k effekt;
        # Qolgani sariq effekt.
        total_teachers = len(raw_leaderboard)
        teacher_leaderboard = []
        for idx, item in enumerate(raw_leaderboard):
            rank = idx + 1
            pct_position = (rank / total_teachers * 100) if total_teachers > 0 else 100

            if rank in [1, 2, 3]:
                tier_class = "bg-gradient-to-r from-amber-100 via-yellow-200 to-amber-100 border-yellow-400 text-yellow-950 ring-2 ring-yellow-400 shadow-md font-bold"
                medal = "🥇" if rank == 1 else ("🥈" if rank == 2 else "🥉")
                badge = f"{medal} {rank}-O'rin (Tilla Lider)"
            elif pct_position <= 10:
                tier_class = "bg-green-100 border-green-400 text-green-900 font-semibold"
                badge = f"🟢 Top 10% (#{rank})"
            elif pct_position <= 70:
                tier_class = "bg-blue-50 border-blue-300 text-blue-900 font-semibold"
                badge = f"🔵 Top 70% (#{rank})"
            else:
                tier_class = "bg-yellow-50 border-yellow-300 text-yellow-800"
                badge = f"🟡 #{rank}"

            item['rank'] = rank
            item['tier_class'] = tier_class
            item['badge'] = badge
            teacher_leaderboard.append(item)

        return render(request, 'accounts/dashboard.html', {
            'my_tests': tests,
            'complaints': Complaint.objects.filter(question__test__creator=request.user).order_by('-created_at'),
            'blocked_attempts': TestAttempt.objects.filter(test__creator=request.user, status='blocked'),
            'teacher_leaderboard': teacher_leaderboard,
        })
    
    student_attempts = TestAttempt.objects.filter(student=request.user, status='completed').exclude(test__test_type__in=['practice', 'anonymous'])
    
    week_ago = now - timezone.timedelta(days=7)
    month_ago = now - timezone.timedelta(days=30)
    weekly_atts = student_attempts.filter(finished_at__gte=week_ago)
    monthly_atts = student_attempts.filter(finished_at__gte=month_ago)

    weekly_score = weekly_atts.aggregate(avg=Avg('score'))['avg'] or 0
    monthly_score = monthly_atts.aggregate(avg=Avg('score'))['avg'] or 0

    weekly_pct = 0
    if weekly_atts.exists():
        pct_list = []
        for a in weekly_atts:
            max_p = sum(q.points for q in a.test.questions.all()) or a.total_questions
            if max_p > 0:
                pct_list.append((a.score / max_p) * 100)
        weekly_pct = sum(pct_list) / len(pct_list) if pct_list else 0

    monthly_pct = 0
    if monthly_atts.exists():
        pct_list = []
        for a in monthly_atts:
            max_p = sum(q.points for q in a.test.questions.all()) or a.total_questions
            if max_p > 0:
                pct_list.append((a.score / max_p) * 100)
        monthly_pct = sum(pct_list) / len(pct_list) if pct_list else 0

    available_tests = Test.objects.filter(status='published', start_time__lte=now).filter(
        Q(deadline__isnull=True) | Q(deadline__gte=now)
    )
    
    valid_tests = []
    for t in available_tests:
        if t.test_type == 'anonymous': continue
        if request.user.role == 'student':
            if t.target_grade != 'barcha' and t.target_grade != request.user.grade: continue
            if t.target_letter != 'barchasi' and t.target_letter != request.user.letter: continue
        valid_tests.append(t.id)

    final_available_tests = Test.objects.filter(id__in=valid_tests).order_by('-created_at')
    anon_tests = Test.objects.filter(status='published', test_type='anonymous')

    return render(request, 'accounts/dashboard.html', {
        'available_tests': final_available_tests,
        'anon_tests': anon_tests,
        'weekly_score': weekly_score,
        'weekly_pct': weekly_pct,
        'monthly_score': monthly_score,
        'monthly_pct': monthly_pct,
        'stats': student_attempts.aggregate(avg_score=Avg('score'), max_score=Max('score'), total_done=Count('id')),
    })

@login_required
def student_rankings_view(request):
    if request.user.is_teacher(): return HttpResponseForbidden()
    now = timezone.now()
    week_ago = now - timezone.timedelta(days=7)
    month_ago = now - timezone.timedelta(days=30)

    rating_period = request.GET.get('period', 'week')
    r_grade = request.GET.get('r_grade', 'barcha')
    r_letter = request.GET.get('r_letter', 'barchasi')

    period_start = week_ago if rating_period == 'week' else month_ago
    
    students_query = User.objects.filter(role='student')
    if r_grade != 'barcha':
        students_query = students_query.filter(grade=r_grade)
    if r_letter != 'barchasi':
        students_query = students_query.filter(letter=r_letter)

    leaderboard_data = []
    for s in students_query:
        atts_in_period = TestAttempt.objects.filter(student=s, status='completed', finished_at__gte=period_start).exclude(test__test_type__in=['practice', 'anonymous'])
        total_tests_done = atts_in_period.count()
        
        avg_pct = 0.0
        if total_tests_done > 0:
            pct_sum = 0.0
            for att in atts_in_period:
                max_p = sum(q.points for q in att.test.questions.all()) or att.total_questions
                if max_p > 0:
                    pct_sum += (att.score / max_p) * 100
            avg_pct = pct_sum / total_tests_done
        
        if total_tests_done > 0 or atts_in_period.exists():
            leaderboard_data.append({
                'user': s,
                'avg_pct': avg_pct,
                'total_tests': total_tests_done
            })

    leaderboard_data = sorted(leaderboard_data, key=lambda x: (x['avg_pct'], x['total_tests']), reverse=True)

    total_count = len(leaderboard_data)
    ranked_list = []
    for idx, item in enumerate(leaderboard_data):
        rank = idx + 1
        pct_position = (rank / total_count * 100) if total_count > 0 else 100
        
        if rank == 1:
            tier_class = "bg-gradient-to-r from-amber-100 to-yellow-200 border-yellow-400 text-yellow-900 ring-4 ring-yellow-400 shadow-xl animate-pulse font-bold"
            badge = "👑 1-O'rin (Chempion)"
        elif rank == 2:
            tier_class = "bg-gradient-to-r from-slate-100 to-gray-200 border-gray-400 text-gray-950 ring-2 ring-gray-400 shadow-lg font-bold"
            badge = "🥈 2-O'rin"
        elif rank == 3:
            tier_class = "bg-gradient-to-r from-orange-50 to-amber-100 border-amber-400 text-amber-950 ring-2 ring-amber-300 shadow-md font-bold"
            badge = "🥉 3-O'rin"
        elif pct_position <= 10:
            tier_class = "bg-green-50 border-green-300 text-green-800 font-semibold"
            badge = f"Top 10% (#{rank})"
        elif pct_position <= 30:
            tier_class = "bg-blue-50 border-blue-300 text-blue-800 font-semibold"
            badge = f"Top 30% (#{rank})"
        elif pct_position <= 70:
            tier_class = "bg-yellow-50 border-yellow-200 text-yellow-800"
            badge = f"Top 70% (#{rank})"
        else:
            tier_class = "bg-red-50 border-red-200 text-red-700"
            badge = f"#{rank}"

        ranked_list.append({
            'rank': rank,
            'user': item['user'],
            'avg_pct': item['avg_pct'],
            'total_tests': item['total_tests'],
            'tier_class': tier_class,
            'badge': badge
        })

    return render(request, 'accounts/student_rankings.html', {
        'ranked_list': ranked_list,
        'rating_period': rating_period,
        'r_grade': r_grade,
        'r_letter': r_letter,
    })

@login_required
def student_results_view(request):
    if request.user.is_teacher(): return HttpResponseForbidden()
    my_attempts = TestAttempt.objects.filter(student=request.user).order_by('-started_at')
    return render(request, 'accounts/student_results.html', {'my_attempts': my_attempts})

@login_required
def teacher_tests_view(request):
    if not request.user.is_teacher(): return HttpResponseForbidden()
    tests = Test.objects.filter(creator=request.user).annotate(attempt_count=Count('attempts'), avg_score=Avg('attempts__score'))
    return render(request, 'accounts/teacher_tests.html', {'my_tests': tests})

@login_required
def teacher_complaints_view(request):
    if not request.user.is_teacher(): return HttpResponseForbidden()
    complaints = Complaint.objects.filter(question__test__creator=request.user).order_by('-created_at')
    return render(request, 'accounts/teacher_complaints.html', {'complaints': complaints})

@login_required
def students_list_view(request):
    if not request.user.is_teacher(): return HttpResponseForbidden()
    students = User.objects.filter(role='student').order_by('grade', 'letter', 'full_name')
    student_data = []
    for s in students:
        last_attempt = TestAttempt.objects.filter(student=s).order_by('-started_at').first()
        student_data.append({
            'student': s,
            'last_attempt': last_attempt
        })
    return render(request, 'accounts/students_list.html', {'student_data': student_data})

@login_required
def close_test(request, test_id):
    test = get_object_or_404(Test, id=test_id)
    if not test.is_owner(request.user): return HttpResponseForbidden()
    test.status = 'closed'
    test.save()
    for att in test.attempts.filter(status='in_progress'):
        att.status = 'completed'
        att.finished_at = timezone.now()
        att.save()
    recalculate_test_scores(test)
    sync_test_to_teacher_google_cloud(test)
    messages.success(request, f"'{test.title}' testi yakunlandi va shaxsiy Google Cloud'da yangilandi!")
    return redirect('teacher_tests')

@login_required
def download_sample_docx(request):
    d = Document()
    d.add_paragraph("EDU-TEST Platformasi - Kengaytirilgan Namuna Test Fayli")
    d.add_paragraph("Eslatma: To'g'ri javob oldiga '*' belgisi qo'yilishi shart. Ochiq savollar uchun 'Javob: ...' yoziladi.")
    
    d.add_paragraph("1. O'zbekiston poytaxti qayer?")
    d.add_paragraph("*A) Toshkent")
    d.add_paragraph("B) Samarqand")
    d.add_paragraph("C) Buxoro")
    d.add_paragraph("D) Xiva")
    
    d.add_paragraph("2. Quyidagilardan qaysilari dasturlash tillari hisoblanadi? (Bir nechta javob)")
    d.add_paragraph("*A) Python")
    d.add_paragraph("B) HTML")
    d.add_paragraph("*C) C++")
    d.add_paragraph("D) Photoshop")

    d.add_paragraph("3. 2 + 2 nechiga teng?")
    d.add_paragraph("Javob: 4, to'rt")
    
    b = io.BytesIO(); d.save(b); b.seek(0)
    res = HttpResponse(b.getvalue(), content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
    res['Content-Disposition'] = 'attachment; filename="namuna_kengaytirilgan_test.docx"'
    return res

@login_required
def create_or_edit_test(request, test_id=None):
    if not request.user.is_teacher(): return HttpResponseForbidden()
    test = get_object_or_404(Test, id=test_id) if test_id else None
    if test and not test.is_owner(request.user): return HttpResponseForbidden()
    
    if request.method == 'POST':
        action_type = request.POST.get('action_type')
        
        if action_type == 'delete_question':
            q_id = request.POST.get('question_id')
            Question.objects.filter(id=q_id, test=test).delete()
            recalculate_test_scores(test)
            sync_test_to_teacher_google_cloud(test)
            return redirect('test_edit', test_id=test.id)

        st = 'draft' if request.POST.get('save_mode') == 'draft' else 'published'
        dur = int(request.POST.get('duration_minutes') or 45)
        dl = request.POST.get('deadline')
        deadline_dt = parse_datetime(dl) if dl else None
        if deadline_dt and timezone.is_naive(deadline_dt): deadline_dt = timezone.make_aware(deadline_dt)
        
        target_g = request.POST.get('target_grade', 'barcha')
        target_l = request.POST.get('target_letter', 'barchasi')
        test_type = request.POST.get('test_type', 'exam')
        allow_t = bool(request.POST.get('allow_teachers'))

        if not test:
            test = Test.objects.create(
                title=str(request.POST.get('title')).strip().capitalize(),
                subject=str(request.POST.get('subject')).strip().capitalize(),
                creator=request.user, duration_minutes=dur, status=st, deadline=deadline_dt,
                target_grade=target_g, target_letter=target_l, test_type=test_type, allow_teachers=allow_t
            )
        else:
            test.title = str(request.POST.get('title')).strip().capitalize()
            test.subject = str(request.POST.get('subject')).strip().capitalize()
            test.duration_minutes = dur
            test.status = st
            test.deadline = deadline_dt
            test.target_grade = target_g
            test.target_letter = target_l
            test.test_type = test_type
            test.allow_teachers = allow_t
            test.save()

        # Fayldan yuklash
        if request.FILES.get('docx_file'): 
            parse_docx_to_test(request.FILES['docx_file'], test)

        # Mavjud savollarning barchasini tahrirlab saqlash
        for q in test.questions.all():
            q_text_val = request.POST.get(f'q_text_{q.id}')
            if q_text_val:
                q.text = str(q_text_val).strip().capitalize()
                try:
                    q.points = float(str(request.POST.get(f'q_points_{q.id}', q.points)).replace(',', '.'))
                except ValueError:
                    pass
                q.correct_text_answers = str(request.POST.get(f'q_open_ans_{q.id}', q.correct_text_answers or '')).strip().capitalize()
                q.save()
                
                corr_opts = request.POST.getlist(f'q_correct_opts_{q.id}')
                for opt in q.options.all():
                    opt_val = request.POST.get(f'opt_text_{opt.id}')
                    if opt_val:
                        opt.text = str(opt_val).strip().capitalize()
                        opt.is_correct = str(opt.id) in corr_opts
                        opt.save()

        # Yangi qo'lda qo'shilgan savol
        qt = request.POST.get('manual_q_text','').strip()
        if qt:
            q_type = request.POST.get('manual_q_type', 'single')
            try:
                q_points = float(str(request.POST.get('manual_q_points', '1')).replace(',', '.'))
            except ValueError:
                q_points = 1.0
            correct_txt = request.POST.get('manual_correct_text', '').strip()
            
            q = Question.objects.create(
                test=test, text=str(qt).strip().capitalize(), question_type=q_type, points=q_points,
                correct_text_answers=str(correct_txt).strip().capitalize(), order=test.questions.count()+1
            )
            if q_type in ['single', 'multiple']:
                correct_indices = request.POST.getlist('manual_correct_opts')
                for i in range(1, 5):
                    ot = request.POST.get(f'manual_opt_{i}','').strip()
                    if ot:
                        is_corr = str(i) in correct_indices
                        Option.objects.create(question=q, text=str(ot).strip().capitalize(), is_correct=is_corr)
        
        recalculate_test_scores(test)
        # 4-talab: Testni to'g'ridan-to'g'ri o'qituvchining o'z Google Cloud xotirasiga saqlash
        sync_test_to_teacher_google_cloud(test)
        messages.success(request, f"Test saqlandi va shaxsiy Google Cloud ({test.teacher_cloud_provider}) xotirangizga yuklandi!")
        return redirect('test_edit', test_id=test.id)
    
    return render(request, 'accounts/test_form.html', {'test': test})

@login_required
def export_test_cloud_json(request, test_id):
    """O'qituvchining Google Cloud'dagi test faylini JSON ko'rinishida yuklab olish yoki ko'rish."""
    test = get_object_or_404(Test, id=test_id)
    if not test.is_owner(request.user):
        return HttpResponseForbidden()
    return JsonResponse(serialize_test_payload(test), json_dumps_params={'ensure_ascii': False, 'indent': 2})

@login_required
def view_or_take_test(request, test_id):
    test = get_object_or_404(Test, id=test_id)
    owner = test.is_owner(request.user)
    if not owner and (test.status == 'draft' or test.status == 'closed' or test.start_time > timezone.now()):
        return HttpResponseForbidden()
    
    if owner: 
        question_stats = []
        for q in test.questions.prefetch_related('options'):
            opt_labels = []
            opt_counts = []
            for opt in q.options.all():
                cnt = StudentAnswer.objects.filter(attempt__test=test, selected_option=opt).count()
                opt_labels.append(opt.text)
                opt_counts.append(cnt)
            question_stats.append({
                'question': q, 'labels': opt_labels, 'counts': opt_counts
            })
        attempts = test.attempts.select_related('student').prefetch_related('answers__question', 'answers__selected_option').order_by('-started_at')
        return render(request, 'accounts/take_test.html', {
            'test': test, 'is_owner_preview': True, 'question_stats': question_stats, 'attempts': attempts
        })
    
    student_obj = None if test.test_type == 'anonymous' else request.user
    
    user_complaints = Complaint.objects.filter(student=request.user, question__test=test) if request.user.is_authenticated else []
    complaint_comments = {c.question_id: c.teacher_comment for c in user_complaints if c.teacher_comment}

    if test.test_type == 'practice':
        last_att = TestAttempt.objects.filter(student=student_obj, test=test).order_by('-started_at').first()
        if not last_att or last_att.status == 'completed':
            att = TestAttempt.objects.create(student=student_obj, test=test, total_questions=test.questions.count(), attempt_number=(last_att.attempt_number + 1 if last_att else 1))
        else:
            att = last_att
    else:
        att, created = TestAttempt.objects.get_or_create(
            student=student_obj, test=test, defaults={'total_questions': test.questions.count()}
        )
        if att.status in ['completed', 'disqualified'] and test.test_type == 'exam':
            return redirect('dashboard')

    return render(request, 'accounts/take_test.html', {'test': test, 'is_owner_preview': False, 'attempt': att, 'complaint_comments': complaint_comments})

@login_required
def submit_test(request, attempt_id):
    att = get_object_or_404(TestAttempt, id=attempt_id)
    if att.student and att.student != request.user: return HttpResponseForbidden()
    
    if request.method == 'POST' and att.status in ['in_progress', 'blocked']:
        att.answers.all().delete()
        score = 0.0; correct = 0
        questions = att.test.questions.prefetch_related('options')

        for q in questions:
            ans_obj = StudentAnswer.objects.create(attempt=att, question=q)
            if q.question_type == 'single':
                sel_id = request.POST.get(f'question_{q.id}')
                if sel_id:
                    opt = q.options.filter(id=sel_id).first()
                    ans_obj.selected_option = opt; ans_obj.save()
                    if opt and opt.is_correct: score += q.points; correct += 1
            elif q.question_type == 'multiple':
                sel_ids = request.POST.getlist(f'question_multi_{q.id}')
                opts = q.options.filter(id__in=sel_ids)
                ans_obj.selected_options.set(opts)
                correct_opts = set(q.options.filter(is_correct=True).values_list('id', flat=True))
                selected_opts = set(opts.values_list('id', flat=True))
                if correct_opts and correct_opts == selected_opts: score += q.points; correct += 1
            elif q.question_type == 'open':
                txt = request.POST.get(f'question_open_{q.id}', '').strip()
                ans_obj.text_answer = str(txt).strip().capitalize(); ans_obj.save()
                if txt and q.correct_text_answers:
                    student_val = str(txt).strip().capitalize()
                    accepted = [str(a).strip().capitalize() for a in q.correct_text_answers.split(',')]
                    if student_val in accepted:
                        score += q.points
                        correct += 1

        att.score = score; att.correct_count = correct; att.total_questions = questions.count()
        att.status = 'completed'; att.finished_at = timezone.now(); att.save()
        messages.success(request, f"Test yakunlandi! Natijangiz: {score} ball.")
    return redirect('dashboard')

@login_required
def report_fullscreen_violation(request, attempt_id):
    att = get_object_or_404(TestAttempt, id=attempt_id)
    if request.method == 'POST' and att.status == 'in_progress':
        att.status = 'blocked'; att.violation_count += 1; att.save()
    return JsonResponse({'status': att.status})

@login_required
def send_excuse_message(request, attempt_id):
    att = get_object_or_404(TestAttempt, id=attempt_id)
    att.excuse_message = request.POST.get('excuse_message','')
    att.save()
    return JsonResponse({'ok': True, 'message': 'Yuborildi!'})

@login_required
def check_attempt_status(request, attempt_id):
    att = get_object_or_404(TestAttempt, id=attempt_id)
    return JsonResponse({'status': att.status, 'remaining_seconds': att.remaining_seconds()})

@login_required
def monitor_students(request, test_id):
    test = get_object_or_404(Test, id=test_id)
    if not test.is_owner(request.user): return HttpResponseForbidden()
    if request.method == 'POST':
        att = get_object_or_404(TestAttempt, id=request.POST.get('attempt_id'), test=test)
        decision = request.POST.get('decision')
        if decision == 'forgive':
            att.status = 'completed' if att.answers.exists() else 'in_progress'
            att.finished_at = timezone.now()
            recalculate_test_scores(test)
            att.save()
        elif decision == 'disqualify':
            att.status = 'disqualified'; att.save()
        elif decision == 'add_retry':
            att.status = 'in_progress'; att.attempt_number += 1; att.save()
        return redirect('monitor_students', test_id=test.id)
    return render(request, 'accounts/monitor.html', {'test': test, 'attempts': test.attempts.select_related('student').all()})

@login_required
def submit_complaint(request, question_id):
    q = get_object_or_404(Question, id=question_id)
    Complaint.objects.create(
        student=request.user if not q.test.test_type == 'anonymous' else None,
        question=q, reason=str(request.POST.get('reason','')).strip().capitalize()
    )
    return JsonResponse({'ok': True, 'message': 'Yuborildi!'})

@login_required
def review_complaint(request, complaint_id):
    c = get_object_or_404(Complaint, id=complaint_id)
    q = c.question
    if not q.test.is_owner(request.user): return HttpResponseForbidden()
    if request.method == 'POST':
        action = request.POST.get('action')
        if action == 'resolve':
            q.text = str(request.POST.get('question_text', q.text)).strip().capitalize()
            try: q.points = float(str(request.POST.get('question_points', q.points)).replace(',', '.'))
            except ValueError: pass
            q.correct_text_answers = str(request.POST.get('correct_text_answers', q.correct_text_answers)).strip().capitalize()
            q.save()
            cid = request.POST.getlist('correct_option_ids')
            for o in q.options.all():
                o.text = str(request.POST.get(f'opt_text_{o.id}', o.text)).strip().capitalize()
                o.is_correct = (str(o.id) in cid); o.save()
            new_opt = request.POST.get('new_option_text', '').strip()
            if new_opt: 
                Option.objects.create(question=q, text=str(new_opt).strip().capitalize(), is_correct=False)
            recalculate_test_scores(q.test)
            sync_test_to_teacher_google_cloud(q.test)
            c.status = 'resolved'
        else:
            c.status = 'rejected'
        c.teacher_comment = str(request.POST.get('teacher_comment', '')).strip().capitalize()
        c.save()
        return redirect('dashboard')
    return render(request, 'accounts/complaint_review.html', {'complaint': c, 'question': q})

# Mobil ilova (API) uchun tayyor JSON endpoint
@login_required
def api_overview(request):
    return JsonResponse({
        "user": {
            "id": request.user.id,
            "full_name": request.user.get_display_name(),
            "role": request.user.role,
            "google_email": request.user.google_email,
            "google_cloud_bucket": request.user.google_cloud_bucket,
        },
        "tests_count": Test.objects.filter(creator=request.user).count() if request.user.is_teacher() else Test.objects.filter(status='published').count()
    })
